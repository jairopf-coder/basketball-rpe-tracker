#!/usr/bin/env python3
"""
verificar_export.py — comprueba un .xlsx de BasketballRPE-Web (Ensino Lugo) contra la
especificación de exportación del fisio y imprime un RESUMEN para comparar con lo importado.

Formato 2: las jugadoras se identifican por nombre y apellido (sin IDs). Columnas de identificación:
«Jugadora» (nombre + apellido, como el «Nombre del Jugador» del CSV de Oli) y «Apellido».

Uso:
    python3 verificar_export.py archivo.xlsx
    python3 verificar_export.py archivo.xlsx --importado-bienestar 42 --importado-rpe 30 --importado-regla 4

Necesita:  pip install openpyxl

Códigos de salida:
    0  el archivo es correcto (y, si se pasaron recuentos importados, coinciden)
    1  el archivo tiene errores (no importarlo hasta corregirlos)
    2  el archivo es correcto pero lo importado NO coincide con lo enviado
"""
import argparse
import datetime
import re
import sys
import unicodedata
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s if s is not None else "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


# canónico -> alias aceptados (sección 5 de la especificación), todo normalizado
ALIAS = {
    "jugadora": ["jugadora", "nombre", "nombre del jugador", "nombre de la jugadora", "player"],
    "apellido": ["apellido", "apellidos", "surname", "last name"],
    "fecha": ["fecha", "date", "dia"],
    "sueno": ["sueno", "calidad del sueno", "sleep"],
    "estres": ["estres", "stress"],
    "fatiga": ["fatiga", "fatigue"],
    "dolor": ["dolor muscular", "doms", "dolor", "muscle soreness"],
    "sesion": ["sesion", "nombre de sesion", "id de sesion", "session"],
    "rpe": ["rpe", "esfuerzo percibido", "srpe", "borg"],
    "activa": ["activa", "regla", "regla activa", "menstruacion", "periodo", "active"],
}
TO_CANON = {a: c for c, al in ALIAS.items() for a in al}
SI = {"si", "1", "verdadero", "true"}
NO = {"no", "0", "falso", "false"}


class Informe:
    def __init__(self):
        self.errores, self.avisos, self.info = [], [], []

    def error(self, m): self.errores.append(m)
    def aviso(self, m): self.avisos.append(m)
    def nota(self, m): self.info.append(m)


def a_fecha(v):
    """Devuelve (date, es_texto_iso) o (None, False)."""
    if isinstance(v, datetime.datetime):
        return v.date(), False
    if isinstance(v, datetime.date):
        return v, False
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v.strip()):
        try:
            return datetime.date.fromisoformat(v.strip()), True
        except ValueError:
            return None, False
    return None, False


def es_entero(v, lo, hi):
    if isinstance(v, bool):
        return False
    if isinstance(v, int):
        return lo <= v <= hi
    if isinstance(v, float) and v.is_integer():
        return lo <= int(v) <= hi
    return False


def leer_hoja(ws, inf, nombre):
    """Cabeceras (fila 1) -> índice de columna; filas de datos como dicts canónicos."""
    if ws.merged_cells.ranges:
        inf.error(f"[{nombre}] hay celdas combinadas.")
    cab = {}
    extra = []
    for i, c in enumerate(ws[1]):
        if c.value is None:
            continue
        k = TO_CANON.get(norm(c.value))
        if k and k not in cab:
            cab[k] = i
        else:
            extra.append(str(c.value))
    filas, vistas_vacias = [], []
    for n, row in enumerate(ws.iter_rows(min_row=2), start=2):
        vals = [c.value for c in row]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            vistas_vacias.append(n)
            continue
        if vistas_vacias:
            inf.error(f"[{nombre}] fila en blanco intercalada en la fila {vistas_vacias[0]}.")
            vistas_vacias = []
        for c in row:
            if c.data_type == "f":
                inf.error(f"[{nombre}] fila {n}: contiene una fórmula (deben ser valores literales).")
                break
        d = {k: (row[i].value if i < len(row) else None) for k, i in cab.items()}
        d["_fila"] = n
        d["_extra"] = {ws.cell(1, j + 1).value: row[j].value for j in range(len(row)) if ws.cell(1, j + 1).value and norm(ws.cell(1, j + 1).value) not in TO_CANON}
        filas.append(d)
    return cab, extra, filas


def clasificar(cab):
    k = set(cab)
    if {"sueno", "estres", "fatiga", "dolor"} <= k:
        return "Bienestar"
    if "activa" in k:
        return "Regla"
    if "rpe" in k:
        return "RPE"
    return None


def comunes(filas, nombre, inf, fechas):
    for f in filas:
        n = f["_fila"]
        if not str(f.get("jugadora") or "").strip():
            inf.error(f"[{nombre}] fila {n}: falta Jugadora.")
        d, texto = a_fecha(f.get("fecha"))
        if d is None:
            inf.error(f"[{nombre}] fila {n}: Fecha no válida ({f.get('fecha')!r}); usa aaaa-mm-dd.")
        else:
            f["_d"] = d
            fechas.append(d)
            if not texto:
                inf.aviso(f"[{nombre}] fila {n}: la fecha es una fecha de Excel, no texto ISO (aceptable).")


def verificar(ruta):
    inf = Informe()
    wb = openpyxl.load_workbook(ruta)  # sin data_only: así se detectan fórmulas
    hojas, leeme = {}, None
    for ws in wb.worksheets:
        cab, extra, filas = leer_hoja(ws, inf, ws.title)
        tipo = clasificar(cab)
        if tipo is None:
            enc = [str(c.value) for c in ws[1] if c.value is not None]
            if [norm(x) for x in enc] == ["clave", "valor"]:
                leeme = {str(r[0].value): r[1].value for r in ws.iter_rows(min_row=2) if r[0].value}
                inf.nota(f"Hoja informativa «{ws.title}» detectada ({len(leeme)} claves).")
            else:
                inf.aviso(f"Hoja «{ws.title}» no reconocida por sus cabeceras: se ignorará.")
            continue
        if tipo in hojas:
            inf.error(f"Hay dos hojas que parecen «{tipo}».")
            continue
        for oblig in ("jugadora", "apellido", "fecha"):
            if oblig not in cab:
                inf.error(f"[{tipo}] falta la columna obligatoria {oblig}.")
        hojas[tipo] = (ws.title, cab, extra, filas)
    for t in ("Bienestar", "RPE", "Regla"):
        if t not in hojas:
            inf.error(f"Falta la hoja «{t}» (se identifica por sus cabeceras).")

    fechas = defaultdict(list)
    ident = defaultdict(set)    # Jugadora -> apellidos con los que aparece (en las tres hojas)

    for tipo, (titulo, cab, extra, filas) in hojas.items():
        comunes(filas, tipo, inf, fechas[tipo])
        for f in filas:
            nom = str(f.get("jugadora") or "").strip()
            ape = str(f.get("apellido") if f.get("apellido") is not None else "").strip()
            ident[nom].add(ape)
            if ape and norm(nom) == norm(ape):
                inf.error(f"[{tipo}] fila {f['_fila']}: Jugadora («{nom}») es igual que Apellido: falta el nombre.")
            elif ape and not norm(nom).endswith(" " + norm(ape)):
                inf.error(f"[{tipo}] fila {f['_fila']}: Jugadora («{nom}») debe ser nombre + apellido y terminar en Apellido («{ape}»).")

        if tipo == "Bienestar":
            vistos = set()
            for f in filas:
                for k, et in (("sueno", "Sueño"), ("estres", "Estrés"), ("fatiga", "Fatiga"), ("dolor", "Dolor muscular")):
                    if not es_entero(f.get(k), 1, 5):
                        inf.error(f"[Bienestar] fila {f['_fila']}: {et} debe ser un entero de 1 a 5 (hay {f.get(k)!r}).")
                clave = (str(f.get("jugadora")).strip(), f.get("_d"))
                if clave in vistos:
                    inf.error(f"[Bienestar] fila {f['_fila']}: duplicado de jugadora+día ({clave[0]}, {clave[1]}).")
                vistos.add(clave)

        elif tipo == "RPE":
            vistos = set()
            for f in filas:
                if not es_entero(f.get("rpe"), 0, 10):
                    inf.error(f"[RPE] fila {f['_fila']}: RPE debe ser un entero de 0 a 10 (hay {f.get('rpe')!r}).")
                ses = str(f.get("sesion") or "").strip()
                clave = (str(f.get("jugadora")).strip(), f.get("_d"), ses)
                if clave in vistos:
                    inf.error(f"[RPE] fila {f['_fila']}: duplicado de jugadora+día+sesión ({clave[0]}, {clave[1]}, «{ses}»).")
                vistos.add(clave)
                # coherencia de la media diaria (columnas extra de BasketballRPE-Web)
                ex = {norm(k): v for k, v in f["_extra"].items()}
                det, nses = ex.get("detalle"), ex.get("n sesiones")
                if det is not None and es_entero(f.get("rpe"), 0, 10):
                    nums = [int(x) for x in re.findall(r"(\d+)\s*(?:·|$)", str(det))]
                    if not nums:
                        inf.aviso(f"[RPE] fila {f['_fila']}: no se pudo leer el Detalle ({det!r}).")
                    else:
                        media = int(sum(nums) / len(nums) + 0.5 + 1e-9)
                        if media != int(f["rpe"]):
                            inf.error(f"[RPE] fila {f['_fila']}: RPE={f['rpe']} no coincide con la media redondeada del Detalle ({det!r} → {media}).")
                        if nses is not None and es_entero(nses, 1, 20) and int(nses) != len(nums):
                            inf.error(f"[RPE] fila {f['_fila']}: Nº sesiones={nses} no coincide con el Detalle ({det!r}).")
            if any(f.get("sesion") for f in filas):
                inf.nota("[RPE] hay filas con Sesión rellena.")
            else:
                inf.nota("[RPE] columna Sesión vacía: el RPE es MEDIA DIARIA por jugadora (convención de BasketballRPE-Web).")
            if any(k for k in ("srpe",) if k in cab):
                inf.aviso("[RPE] la columna RPE se llama sRPE/alias de s-RPE: el s-RPE NO debe exportarse.")

        elif tipo == "Regla":
            vistos = set()
            por_jug = defaultdict(list)
            valores = set()
            for f in filas:
                v = norm(f.get("activa"))
                if v in SI:
                    marca = True
                elif v in NO:
                    marca = False
                else:
                    inf.error(f"[Regla] fila {f['_fila']}: valor de Activa no reconocido ({f.get('activa')!r}).")
                    continue
                valores.add(marca)
                clave = (str(f.get("jugadora")).strip(), f.get("_d"))
                if clave in vistos:
                    inf.error(f"[Regla] fila {f['_fila']}: duplicado de jugadora+día ({clave[0]}, {clave[1]}).")
                vistos.add(clave)
                if marca and f.get("_d"):
                    por_jug[clave[0]].append(f["_d"])
            if valores == {True}:
                inf.nota("[Regla] solo hay filas «Sí» (convención de BasketballRPE-Web): un día sin fila = sin regla.")
            elif False in valores:
                inf.nota("[Regla] hay filas «No» (no es la convención habitual de BasketballRPE-Web).")
            hasta = None
            if leeme and leeme.get("Ventana hasta"):
                hasta, _ = a_fecha(str(leeme["Ventana hasta"]))
            periodos = {}
            for nom, ds in por_jug.items():
                ds = sorted(set(ds))
                grupos, ini = [], ds[0]
                for a, b in zip(ds, ds[1:] + [None]):
                    if b is None or (b - a).days != 1:
                        grupos.append((ini, a))
                        ini = b
                periodos[nom] = grupos
            hojas["Regla"] = hojas["Regla"] + (periodos, hasta)

    # --- identificación coherente entre hojas (nombre + apellido, sin IDs) ---
    for nom, conj in ident.items():
        if len(conj) > 1:
            inf.error(f"«{nom}» aparece con apellidos distintos: {sorted(conj)} (deben ser idénticos en las tres hojas; "
                      f"si son dos jugadoras con el mismo nombre, no se pueden distinguir).")
    sin_apellido = sorted(n for n, c in ident.items() if c == {""})
    if sin_apellido:
        inf.aviso(f"{len(sin_apellido)} jugadora(s) con Apellido vacío (la ficha aún no lo tiene; se identifican solo por el nombre): "
                  f"{', '.join(sin_apellido)}.")

    # --- recuentos ---
    n = {t: len(hojas[t][3]) for t in hojas}
    todas = [d for t in fechas.values() for d in t]
    desde, hasta_dato = (min(todas), max(todas)) if todas else (None, None)
    jug = len(ident)

    # --- contraste con la hoja LEEME (control interno) ---
    if leeme:
        m = re.search(r"versi[oó]n\s+(\d+)", str(leeme.get("Formato", "")))
        if m and int(m.group(1)) != 2:
            inf.error(f"El archivo es de formato {m.group(1)}; este verificador espera el formato 2 (identificación por nombre y apellido).")
        elif not m:
            inf.aviso("LEEME no declara la versión del formato (se esperaba «versión 2»).")
        for clave, t in (("Filas Bienestar", "Bienestar"), ("Filas RPE", "RPE"), ("Filas Regla", "Regla")):
            if t in n and str(leeme.get(clave)) != str(n[t]):
                inf.error(f"LEEME dice {clave}={leeme.get(clave)} pero la hoja tiene {n[t]} filas (¿archivo truncado o modificado?).")
        vd, _ = a_fecha(str(leeme.get("Ventana desde", "")))
        vh, _ = a_fecha(str(leeme.get("Ventana hasta", "")))
        if vd and vh:
            fuera = [d for d in todas if d < vd or d > vh]
            if fuera:
                inf.error(f"{len(fuera)} fila(s) con fecha fuera de la ventana {vd} a {vh} declarada en LEEME.")
        else:
            inf.aviso("LEEME no trae una ventana desde/hasta legible.")
    return inf, hojas, n, jug, desde, hasta_dato


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo")
    ap.add_argument("--importado-bienestar", type=int)
    ap.add_argument("--importado-rpe", type=int)
    ap.add_argument("--importado-regla", type=int)
    a = ap.parse_args()

    inf, hojas, n, jug, desde, hasta = verificar(a.archivo)

    print(f"\nArchivo: {a.archivo}")
    print("-" * 60)
    for t in ("Bienestar", "RPE", "Regla"):
        if t in n:
            print(f"  {t:<10} {n[t]:>5} filas   (hoja «{hojas[t][0]}»)")
    print(f"  Jugadoras distintas: {jug}")
    if desde:
        print(f"  Fechas en los datos: {desde} a {hasta}")

    if "Regla" in hojas and len(hojas["Regla"]) == 6:
        periodos = hojas["Regla"][4]
        ventana_hasta = hojas["Regla"][5]
        if periodos:
            print("\nPeríodos de regla que debería reconstruir el importador (días consecutivos con «Sí»):")
            for nom in sorted(periodos):
                for ini, fin in periodos[nom]:
                    abierto = " (ABIERTO: llega al último día de la ventana)" if ventana_hasta and fin >= ventana_hasta else ""
                    print(f"  - {nom}: {ini} a {fin}{abierto}")

    for titulo, lista, marca in (("ERRORES", inf.errores, "✗"), ("AVISOS", inf.avisos, "!"), ("NOTAS", inf.info, "i")):
        if lista:
            print(f"\n{titulo}:")
            for m in lista[:60]:
                print(f"  {marca} {m}")
            if len(lista) > 60:
                print(f"  … y {len(lista) - 60} más")

    resumen = f"RESUMEN: Bienestar={n.get('Bienestar', 0)} RPE={n.get('RPE', 0)} Regla={n.get('Regla', 0)} Jugadoras={jug} Desde={desde} Hasta={hasta}"
    print("\n" + resumen)

    codigo = 1 if inf.errores else 0
    esperados = {"Bienestar": a.importado_bienestar, "RPE": a.importado_rpe, "Regla": a.importado_regla}
    if any(v is not None for v in esperados.values()):
        print("\nContraste con lo IMPORTADO por la app del estudio:")
        for t, v in esperados.items():
            if v is None:
                continue
            enviado = n.get(t, 0)
            if v == enviado:
                print(f"  ✓ {t}: enviado {enviado} = importado {v}")
            else:
                print(f"  ✗ {t}: enviado {enviado} ≠ importado {v} (diferencia {enviado - v:+d}). Revisar qué filas se descartaron y por qué.")
                if codigo == 0:
                    codigo = 2
    print("\n" + ("RESULTADO: CORRECTO" if codigo == 0 else "RESULTADO: HAY PROBLEMAS (código %d)" % codigo))
    sys.exit(codigo)


if __name__ == "__main__":
    main()
